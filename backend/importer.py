from __future__ import annotations

import csv
from io import BytesIO, StringIO

from openpyxl import load_workbook


HEADER_ALIASES = {
  "device_name": "name",
  "devicename": "name",
  "name": "name",
  "device": "name",
  "type": "type",
  "category": "type",
  "device_type": "type",
  "ip": "ipAddress",
  "ip_address": "ipAddress",
  "ipaddress": "ipAddress",
  "ipaddr": "ipAddress",
  "ip_add": "ipAddress",
  "location": "location",
  "loc": "location",
}


def _normalize_header(value: object) -> str:
  return str(value or "").strip().lower().replace(" ", "_")


def _map_row(headers: list[str], values: list[object]) -> dict:
  mapped = {}
  for header, value in zip(headers, values):
    key = HEADER_ALIASES.get(_normalize_header(header))
    if key:
      mapped[key] = str(value or "").strip()
  return mapped


def parse_csv(content: bytes) -> list[dict]:
  text = content.decode("utf-8-sig")
  reader = csv.reader(StringIO(text))
  rows = list(reader)
  if not rows:
    return []
  headers = rows[0]
  return [_map_row(headers, row) for row in rows[1:] if any(str(cell or "").strip() for cell in row)]


def parse_xlsx(content: bytes) -> list[dict]:
  workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
  sheet = workbook.active
  rows = list(sheet.iter_rows(values_only=True))
  if not rows:
    return []
  headers = [str(value or "") for value in rows[0]]
  return [_map_row(headers, list(row)) for row in rows[1:] if any(str(cell or "").strip() for cell in row)]


def parse_device_file(filename: str, content: bytes) -> list[dict]:
  lower = filename.lower()
  if lower.endswith(".csv"):
    return parse_csv(content)
  if lower.endswith(".xlsx"):
    return parse_xlsx(content)
  raise ValueError("Upload .xlsx or .csv file only.")
